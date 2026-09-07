import os, json, secrets, time, io, csv, re
from pathlib import Path
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file, abort
from docx import Document
from pypdf import PdfReader
from db import connect, init_db
from scoring import score_text, proofread_text, DIMS, AREAS

app=Flask(__name__)
app.config["SECRET_KEY"]=os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config["MAX_CONTENT_LENGTH"]=20*1024*1024
app.config["SESSION_COOKIE_HTTPONLY"]=True
app.config["SESSION_COOKIE_SAMESITE"]="Lax"
app.config["SESSION_COOKIE_SECURE"]=os.environ.get("COOKIE_SECURE","0")=="1"

MODELS=[
    {"id":"target","name":"待测政策文本"},
    {"id":"deepseek","name":"DeepSeek"},
    {"id":"kimi","name":"Kimi"},
    {"id":"glm53","name":"GLM5.3"},
    {"id":"qwen37","name":"Qwen3.7"},
    {"id":"hy3","name":"HY3"},
]
ALLOWED={".txt",".md",".docx",".pdf",".html",".htm"}
RATE={}

def limited(bucket,max_count=40,window=60):
    ip=request.headers.get("X-Forwarded-For",request.remote_addr or "").split(",")[0].strip()
    key=(bucket,ip); now=time.time(); arr=[t for t in RATE.get(key,[]) if now-t<window]
    if len(arr)>=max_count:return True
    arr.append(now); RATE[key]=arr; return False

def csrf_token():
    if "_csrf" not in session: session["_csrf"]=secrets.token_urlsafe(24)
    return session["_csrf"]
app.jinja_env.globals["csrf_token"]=csrf_token

def require_csrf():
    token=request.form.get("_csrf") or request.headers.get("X-CSRF-Token")
    if not token or token!=session.get("_csrf"): abort(400,"CSRF validation failed")

def extract_upload(file):
    if not file or not file.filename:return ""
    ext=Path(file.filename).suffix.lower()
    if ext not in ALLOWED: raise ValueError("不支持的文件格式")
    data=file.read()
    if len(data)>3*1024*1024: raise ValueError("单个文件不能超过3MB")
    if ext in {".txt",".md"}: return data.decode("utf-8",errors="replace")
    if ext in {".html",".htm"}:
        text=data.decode("utf-8",errors="replace")
        return re.sub(r"\s+"," ",re.sub(r"<[^>]+>"," ",text)).strip()
    if ext==".docx":
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
            tmp.write(data); tmp.flush(); doc=Document(tmp.name)
            return "\n".join(p.text for p in doc.paragraphs)
    if ext==".pdf":
        reader=PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    return ""

def get_assessment(token):
    with connect() as db: row=db.execute("SELECT * FROM assessments WHERE access_token=?",(token,)).fetchone()
    if not row: abort(404)
    d=dict(row)
    for k in ["areas","inputs","results","reviews","proofreading"]: d[k]=json.loads(d[k] or ("[]" if k=="areas" else "{}"))
    return d

def final_score(x,review,d): return review.get("scores",{}).get(d["key"],x["score"][d["key"]])
def final_total(mid,x,reviews): return round(sum(float(final_score(x,reviews.get(mid,{}),d)) for d in DIMS),1)


def get_version_assessment(token):
    with connect() as db:
        row=db.execute("SELECT * FROM version_assessments WHERE access_token=?",(token,)).fetchone()
    if not row: abort(404)
    d=dict(row)
    for k in ["areas","previous_score","current_score","current_proofreading"]:
        d[k]=json.loads(d[k] or ("[]" if k=="areas" else "{}"))
    return d

def score_delta(previous,current):
    return {d["key"]: round(float(current.get(d["key"],0))-float(previous.get(d["key"],0)),1) for d in DIMS}

@app.after_request
def security_headers(resp):
    resp.headers["X-Content-Type-Options"]="nosniff"
    resp.headers["X-Frame-Options"]="DENY"
    resp.headers["Referrer-Policy"]="same-origin"
    resp.headers["Permissions-Policy"]="camera=(), microphone=(), geolocation=()"
    resp.headers["Content-Security-Policy"]="default-src 'self'; style-src 'self'; script-src 'self'; img-src 'self' data:; form-action 'self'; frame-ancestors 'none'"
    return resp

@app.route("/")
def index():
    recent=[]
    tokens=session.get("recent",[])[:8]
    if tokens:
        with connect() as db:
            for t in tokens:
                r=db.execute("SELECT access_token,code,title,created_at FROM assessments WHERE access_token=?",(t,)).fetchone()
                if r: recent.append(r)
    recent_versions=[]
    version_tokens=session.get("recent_versions",[])[:8]
    if version_tokens:
        with connect() as db:
            for t in version_tokens:
                r=db.execute("SELECT access_token,code,title,created_at FROM version_assessments WHERE access_token=?",(t,)).fetchone()
                if r: recent_versions.append(r)
    return render_template("dashboard.html",models=MODELS,areas=AREAS,recent=recent,recent_versions=recent_versions)

@app.post("/assessment")
def create_assessment():
    if limited("assessment",20,60): abort(429)
    require_csrf()
    title=request.form.get("title","").strip()
    code=request.form.get("code","").strip() or f"POL-{int(time.time())}"
    areas=request.form.getlist("areas") or ["general"]
    inputs={}; results={}; proofreading={}
    for m in MODELS:
        text=request.form.get("text_"+m["id"],"").strip()
        file=request.files.get("file_"+m["id"])
        if file and file.filename:
            try:
                uploaded=extract_upload(file).strip()
                if uploaded:text=uploaded
            except ValueError as e:
                flash(f'{m["name"]}：{e}',"error"); return redirect(url_for("index"))
        if text:
            inputs[m["id"]]=text
            results[m["id"]]={**m,"score":score_text(text,areas)}
            if m["id"]=="target": proofreading[m["id"]]=proofread_text(text)
    if not results:
        flash("请至少输入或上传一份待评价文本。","error"); return redirect(url_for("index"))
    token=secrets.token_urlsafe(32)
    with connect() as db:
        db.execute("""INSERT INTO assessments(access_token,code,title,areas,inputs,results,proofreading)
                      VALUES(?,?,?,?,?,?,?)""",
                   (token,code,title,json.dumps(areas,ensure_ascii=False),json.dumps(inputs,ensure_ascii=False),json.dumps(results,ensure_ascii=False),json.dumps(proofreading,ensure_ascii=False)))
    recent=[token]+[t for t in session.get("recent",[]) if t!=token]
    session["recent"]=recent[:8]
    return redirect(url_for("assessment",token=token))


@app.post("/version-assessment")
def create_version_assessment():
    if limited("version_assessment",20,60): abort(429)
    require_csrf()
    title=request.form.get("version_title","").strip()
    code=request.form.get("version_code","").strip() or f"VER-{int(time.time())}"
    areas=request.form.getlist("version_areas") or ["general"]

    def load_text(field):
        text=request.form.get(field,"").strip()
        file=request.files.get("file_"+field)
        if file and file.filename:
            uploaded=extract_upload(file).strip()
            if uploaded: text=uploaded
        return text

    try:
        previous_text=load_text("previous_text")
        current_text=load_text("current_text")
    except ValueError as e:
        flash(str(e),"error"); return redirect(url_for("index")+"#version-assessment")
    if not previous_text or not current_text:
        flash("纵向测评需要同时提供上一版本和当前版本文本。","error")
        return redirect(url_for("index")+"#version-assessment")

    previous_score=score_text(previous_text,areas)
    current_score=score_text(current_text,areas)
    current_proofreading=proofread_text(current_text)
    token=secrets.token_urlsafe(32)
    with connect() as db:
        db.execute("""INSERT INTO version_assessments(access_token,code,title,areas,previous_text,current_text,previous_score,current_score,current_proofreading)
                      VALUES(?,?,?,?,?,?,?,?,?)""",
                   (token,code,title,json.dumps(areas,ensure_ascii=False),previous_text,current_text,
                    json.dumps(previous_score,ensure_ascii=False),json.dumps(current_score,ensure_ascii=False),json.dumps(current_proofreading,ensure_ascii=False)))
    recent=[token]+[t for t in session.get("recent_versions",[]) if t!=token]
    session["recent_versions"]=recent[:8]
    return redirect(url_for("version_result",token=token))

@app.route("/v/<token>")
def version_result(token):
    v=get_version_assessment(token)
    delta=score_delta(v["previous_score"],v["current_score"])
    total_delta=round(float(v["current_score"].get("total",0))-float(v["previous_score"].get("total",0)),1)
    return render_template("longitudinal.html",v=v,dims=DIMS,areas=AREAS,delta=delta,total_delta=total_delta)

@app.route("/v/<token>/csv")
def version_csv(token):
    v=get_version_assessment(token); delta=score_delta(v["previous_score"],v["current_score"])
    sio=io.StringIO(); w=csv.writer(sio)
    w.writerow(["评价维度","上一版本","当前版本","变化值"])
    for d in DIMS:
        w.writerow([d["name"],v["previous_score"].get(d["key"],0),v["current_score"].get(d["key"],0),delta[d["key"]]])
    w.writerow(["综合得分",v["previous_score"].get("total",0),v["current_score"].get("total",0),round(v["current_score"].get("total",0)-v["previous_score"].get("total",0),1)])
    data=("\ufeff"+sio.getvalue()).encode("utf-8")
    return send_file(io.BytesIO(data),mimetype="text/csv",as_attachment=True,download_name=f'{v["code"]}_版本对比结果.csv')

@app.route("/v/<token>/word")
def version_word(token):
    v=get_version_assessment(token); delta=score_delta(v["previous_score"],v["current_score"])
    doc=Document(); doc.add_heading("政策文本版本迭代测评报告",0)
    doc.add_paragraph(f'测评编号：{v["code"]}'); doc.add_paragraph(f'测评任务：{v["title"]}')
    doc.add_heading("一、版本变化概览",1)
    doc.add_paragraph(f'上一版本综合得分：{v["previous_score"].get("total",0)}')
    doc.add_paragraph(f'当前版本综合得分：{v["current_score"].get("total",0)}')
    td=round(v["current_score"].get("total",0)-v["previous_score"].get("total",0),1)
    doc.add_paragraph(f'综合变化：{td:+.1f}分')
    table=doc.add_table(rows=1,cols=4); table.style="Table Grid"
    for i,t in enumerate(["评价维度","上一版本","当前版本","变化"]): table.rows[0].cells[i].text=t
    for d in DIMS:
        cells=table.add_row().cells
        vals=[d["name"],str(v["previous_score"].get(d["key"],0)),str(v["current_score"].get(d["key"],0)),f'{delta[d["key"]]:+.1f}']
        for i,t in enumerate(vals): cells[i].text=t
    doc.add_heading("二、当前版本主要不足",1)
    for item in v["current_score"].get("defects",[]):
        doc.add_paragraph(f'[{item["severity"]}] {item["dimension"]}：{item["issue"]}')
        doc.add_paragraph(f'发现：{item["evidence"]}'); doc.add_paragraph(f'建议：{item["action"]}')
    proof=v.get("current_proofreading",{})
    if proof.get("issues"):
        doc.add_heading("三、当前版本原文审校",1)
        for it in proof["issues"]:
            doc.add_paragraph(f'批注{it["id"]}｜{it["category"]}｜{it["severity"]}风险')
            doc.add_paragraph(f'原文：{it["original"]}')
            doc.add_paragraph(f'问题：{it["problem"]}')
            doc.add_paragraph(f'修改建议：{it["suggestion"]}')
            doc.add_paragraph(f'建议改写：{it["rewrite"]}')
    doc.add_heading("四、方法说明",1)
    doc.add_paragraph("纵向测评对同一政策文本的上一版本与当前版本采用同一套八维规则进行评价，以维度变化和综合变化反映文本迭代效果；自动结果用于辅助分析，不替代政策事实核验和专业判断。")
    bio=io.BytesIO(); doc.save(bio); bio.seek(0)
    return send_file(bio,as_attachment=True,download_name=f'{v["code"]}_版本迭代测评报告.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.route("/a/<token>")
def assessment(token):
    a=get_assessment(token)
    ranks=sorted([(final_total(mid,x,a["reviews"]),mid,x) for mid,x in a["results"].items()],reverse=True,key=lambda t:t[0])
    return render_template("assessment.html",a=a,dims=DIMS,ranks=ranks,areas=AREAS)

@app.post("/a/<token>/review")
def review(token):
    if limited("review",40,60): abort(429)
    require_csrf(); a=get_assessment(token)
    mid=request.form.get("model_id","")
    if mid not in a["results"]: abort(400)
    scores={}; reasons={}
    for d in DIMS:
        raw=request.form.get("score_"+d["key"],"").strip()
        if raw:
            v=max(0,min(float(raw),d["max"])); scores[d["key"]]=v
            reasons[d["key"]]=request.form.get("reason_"+d["key"],"").strip()
    entry={"scores":scores,"reasons":reasons,"reviewer":request.form.get("reviewer","").strip() or "未署名复核人","organization":request.form.get("organization","").strip(),"reviewer_type":request.form.get("reviewer_type","同行/研究人员"),"note":request.form.get("note","").strip(),"time":time.strftime("%Y-%m-%d %H:%M:%S")}
    reviews=a["reviews"]; history=reviews.get(mid,{}).get("history",[]); history.append(entry.copy()); entry["history"]=history; reviews[mid]=entry
    with connect() as db: db.execute("UPDATE assessments SET reviews=?,updated_at=CURRENT_TIMESTAMP WHERE access_token=?",(json.dumps(reviews,ensure_ascii=False),token))
    flash("人工复核已保存。","ok")
    return redirect(url_for("assessment",token=token)+"#review")


@app.route("/a/<token>/export")
def export_page(token):
    a=get_assessment(token)
    ranks=sorted([(final_total(mid,x,a["reviews"]),mid,x) for mid,x in a["results"].items()],reverse=True,key=lambda t:t[0])
    return render_template("export.html",a=a,dims=DIMS,ranks=ranks,areas=AREAS)

@app.route("/a/<token>/proofreading.csv")
def export_proof_csv(token):
    a=get_assessment(token)
    proof=a.get("proofreading",{}).get("target",{})
    sio=io.StringIO(); w=csv.writer(sio)
    w.writerow(["编号","问题类型","风险等级","原文","问题说明","修改建议","建议改写"])
    for it in proof.get("issues",[]):
        w.writerow([it.get("id"),it.get("category"),it.get("severity"),it.get("original"),it.get("problem"),it.get("suggestion"),it.get("rewrite")])
    data=("\ufeff"+sio.getvalue()).encode("utf-8")
    return send_file(io.BytesIO(data),mimetype="text/csv",as_attachment=True,download_name=f'{a["code"]}_政策文本审校清单.csv')

@app.route("/a/<token>/csv")
def export_csv(token):
    a=get_assessment(token); sio=io.StringIO(); w=csv.writer(sio)
    w.writerow(["模型","机器初评","人工复核后总分"]+[d["name"] for d in DIMS])
    for mid,x in a["results"].items():
        vals=[final_score(x,a["reviews"].get(mid,{}),d) for d in DIMS]
        w.writerow([x["name"],x["score"]["total"],round(sum(map(float,vals)),1)]+vals)
    data=("\ufeff"+sio.getvalue()).encode("utf-8")
    return send_file(io.BytesIO(data),mimetype="text/csv",as_attachment=True,download_name=f'{a["code"]}_测评结果.csv')

@app.route("/a/<token>/word")
def export_word(token):
    a=get_assessment(token); doc=Document(); doc.add_heading("AI政策文本测评报告",0)
    doc.add_paragraph(f'测评编号：{a["code"]}'); doc.add_paragraph(f'测评任务：{a["title"]}')
    doc.add_heading("一、横向比较",1)
    table=doc.add_table(rows=1,cols=3); table.style="Table Grid"
    for i,t in enumerate(["评价对象","机器初评","复核后总分"]): table.rows[0].cells[i].text=t
    ranks=sorted([(final_total(mid,x,a["reviews"]),mid,x) for mid,x in a["results"].items()],reverse=True,key=lambda t:t[0])
    for total,mid,x in ranks:
        cells=table.add_row().cells; cells[0].text=x["name"]; cells[1].text=str(x["score"]["total"]); cells[2].text=str(total)
    if "target" in a["results"]:
        doc.add_heading("二、待测政策文本主要不足",1)
        for d in a["results"]["target"]["score"].get("defects",[]):
            doc.add_paragraph(f'[{d["severity"]}] {d["dimension"]}：{d["issue"]}')
            doc.add_paragraph(f'证据：{d["evidence"]}'); doc.add_paragraph(f'建议：{d["action"]}')
        proof=a["proofreading"].get("target",{})
        if proof.get("issues"):
            doc.add_heading("三、原文审校建议",1)
            for it in proof["issues"]:
                doc.add_paragraph(f'批注{it["id"]}｜{it["category"]}｜{it["severity"]}风险')
                doc.add_paragraph(f'原文：{it["original"]}')
                doc.add_paragraph(f'问题：{it["problem"]}')
                doc.add_paragraph(f'修改建议：{it["suggestion"]}')
                doc.add_paragraph(f'建议改写：{it["rewrite"]}')
    doc.add_heading("四、方法说明",1); doc.add_paragraph("本系统采用八维100分制进行结构化初评。自动评价用于发现文本中的可观察问题，不替代政策事实核验和专家判断；人工复核结果单独保留。")
    bio=io.BytesIO(); doc.save(bio); bio.seek(0)
    return send_file(bio,as_attachment=True,download_name=f'{a["code"]}_测评报告.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.route("/method")
def method(): return render_template("method.html",dims=DIMS,areas=AREAS)
@app.route("/help")
def help_page(): return render_template("help.html")

@app.errorhandler(413)
def too_large(e): return "上传内容超过限制",413

with app.app_context(): init_db()
if __name__=="__main__": app.run(host="0.0.0.0",port=8000,debug=False)
